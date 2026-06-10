from __future__ import annotations

import csv
import json
import re
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.recommender.deepseek import normalize_text, rerank_with_deepseek
from app.recommender.models import KeywordRow, Recommendation, TagTaxonomy, VoiceRecord

DEFAULT_RULES: dict[str, Any] = {
    "language_hints": [],
    "ambiguous_keywords": [],
    "attribute_tags": {"child": "", "dialect": "", "multi_language": ""},
    "attribute_hints": {"child": [], "dialect": []},
    "intent_hints": {"formal": [], "service": [], "marketing": []},
    "service_scene_tags": [],
    "marketing_focus": {"l1_tags": [], "l2_tags": [], "short_video_tags": []},
    "marketing_backfill": [],
    "formal_backfill": [],
    "scene_keyword_overrides": [],
    "audio_preferred_words": [],
    "audio_deprioritized_words": [],
    "speaker_dedupe_suffix_patterns": [],
    "speaker_penalties": [],
    "price_pattern": r"\d+(\.\d+)?",
    "tech_freshness_excluded_l1": [],
    "tech_rank_patterns": [],
}

LATIN_LANGUAGE_MARKERS: list[tuple[str, tuple[str, ...]]] = [
    ("越南语", ("ă", "â", "đ", "ê", "ô", "ơ", "ư", "ạ", "ả", "ấ", "ầ", "ẩ", "ẫ", "ậ", "ắ", "ằ", "ẳ", "ẵ", "ặ", "ẹ", "ẻ", "ẽ", "ế", "ề", "ể", "ễ", "ệ", "ị", "ọ", "ỏ", "ố", "ồ", "ổ", "ỗ", "ộ", "ớ", "ờ", "ở", "ỡ", "ợ", "ụ", "ủ", "ứ", "ừ", "ử", "ữ", "ự", "ỳ", "ỵ", "ỷ", "ỹ")),
    ("波兰语", ("ą", "ć", "ę", "ł", "ń", "ó", "ś", "ź", "ż")),
    ("德语", ("ä", "ö", "ü", "ß")),
    ("西班牙语", ("ñ", "¿", "¡")),
    ("葡萄牙语", ("ã", "õ", "ç")),
    ("法语", ("œ", "æ")),
    ("意大利语", (" è ", " é ", " che ", " gli ", " per ", " una ")),
    ("印尼语", (" yang ", " tidak ", " dengan ", " untuk ", " saya ")),
    ("马来语", (" yang ", " tidak ", " dengan ", " untuk ", " saya ")),
    ("菲律宾语", (" ang ", " mga ", " hindi ", " para ", " ako ")),
]

TECH_SUFFIX_PATTERNS = [
    r"[-－—_]\s*超拟人\s*[（(]\s*Pro\s*[）)]\s*$",
    r"[-－—_]\s*超拟人\s*$",
    r"[-－—_]\s*Pro\s*$",
    r"[-－—_]\s*品质\s*$",
    r"[-－—_]\s*默认\s*$",
]


def load_rule_settings(path: Path | None) -> dict[str, Any]:
    if not path or not path.exists():
        return DEFAULT_RULES.copy()
    with path.open("r", encoding="utf-8") as file:
        loaded = json.load(file)
    rules = DEFAULT_RULES.copy()
    for key, value in loaded.items():
        if isinstance(value, dict) and isinstance(rules.get(key), dict):
            merged = dict(rules[key])
            merged.update(value)
            rules[key] = merged
        else:
            rules[key] = value
    return rules


def split_tags(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, list):
        tags: list[str] = []
        for item in value:
            tags.extend(split_tags(item))
        return tags
    parts = re.split(r"[,，、/|]+", str(value).strip())
    return [part.strip() for part in parts if part and part.strip()]


def keyword_weight(keyword: str, rules: dict[str, Any]) -> float:
    value = (keyword or "").strip()
    if not value:
        return 0.0
    if value in set(rules.get("ambiguous_keywords", [])):
        return 0.15
    if len(value) <= 1:
        return 0.1
    if len(value) == 2:
        return 0.45
    return 1.0


def to_gender_label(code: str) -> str:
    if str(code).strip() == "1":
        return "男"
    if str(code).strip() == "2":
        return "女"
    return "未知"


def char_count(pattern: str, text: str) -> int:
    return len(re.findall(pattern, text))


def detect_script_language(text: str, available_languages: set[str]) -> str:
    normalized = f" {normalize_text(text).lower()} "
    if not normalized.strip():
        return ""

    checks = [
        ("日语", r"[\u3040-\u30ff]"),
        ("韩语", r"[\uac00-\ud7af]"),
        ("俄语", r"[\u0400-\u04ff]"),
        ("阿拉伯语", r"[\u0600-\u06ff]"),
        ("泰语", r"[\u0e00-\u0e7f]"),
    ]
    for language, pattern in checks:
        if language in available_languages and char_count(pattern, normalized) >= 2:
            return language

    latin_count = char_count(r"[a-zA-Z]", normalized)
    chinese_count = char_count(r"[\u4e00-\u9fff]", normalized)
    if latin_count < 12 or latin_count <= chinese_count * 1.5:
        return ""

    for language, markers in LATIN_LANGUAGE_MARKERS:
        if language in available_languages and any(marker in normalized for marker in markers):
            return language
    return "英语" if "英语" in available_languages else ""


def load_tag_taxonomy(xlsx_path: Path) -> TagTaxonomy:
    workbook = load_workbook(xlsx_path, data_only=True)
    rows = list(workbook[workbook.sheetnames[0]].iter_rows(values_only=True))
    if not rows:
        raise ValueError("标签体系为空")

    col_to_l1: dict[int, str] = {}
    l1_tags: set[str] = set()
    l1_to_l2: dict[str, set[str]] = {}
    for index, cell in enumerate(rows[0]):
        if index == 0:
            continue
        tag = str(cell or "").strip()
        if not tag:
            continue
        col_to_l1[index] = tag
        l1_tags.add(tag)
        l1_to_l2.setdefault(tag, set())

    l2_tags: set[str] = set()
    for row in rows[1:]:
        for index, l1 in col_to_l1.items():
            value = str((row[index] if index < len(row) else "") or "").strip()
            if value:
                l1_to_l2[l1].add(value)
                l2_tags.add(value)

    l2_to_l1 = {l2: l1 for l1, l2s in l1_to_l2.items() for l2 in l2s}
    return TagTaxonomy(l1_tags=l1_tags, l2_tags=l2_tags, l1_to_l2=l1_to_l2, l2_to_l1=l2_to_l1)


def load_voice_records(xlsx_path: Path, taxonomy: TagTaxonomy) -> list[VoiceRecord]:
    workbook = load_workbook(xlsx_path, data_only=True)
    rows = list(workbook[workbook.sheetnames[0]].iter_rows(values_only=True))
    headers = [str(item).strip() if item else "" for item in rows[0]]
    idx = {header: index for index, header in enumerate(headers) if header}
    required = ["发音人名称（旧）", "VCN", "性别", "语言", "一级场景标签", "二级场景标签", "技术描述（X系列&一句话）"]
    missing = [field for field in required if field not in idx]
    if missing:
        raise ValueError(f"音库缺少必要字段: {missing}")

    records: list[VoiceRecord] = []
    for row in rows[1:]:
        speaker = str(row[idx["发音人名称（旧）"]] or "").strip()
        vcn = str(row[idx["VCN"]] or "").strip()
        if not speaker or not vcn:
            continue
        raw_l1 = split_tags(row[idx["一级场景标签"]])
        raw_l2 = split_tags(row[idx["二级场景标签"]])
        scene_l1 = [tag for tag in raw_l1 if tag in taxonomy.l1_tags] or raw_l1
        scene_l2 = [tag for tag in raw_l2 if tag in taxonomy.l2_tags] or raw_l2
        try:
            sort_value = float(row[idx["排序值"]]) if "排序值" in idx and row[idx["排序值"]] not in (None, "") else 999999.0
        except Exception:
            sort_value = 999999.0
        gender_code = str(row[idx["性别"]] or "").strip()
        records.append(
            VoiceRecord(
                speaker_name=speaker,
                vcn=vcn,
                gender_code=gender_code,
                gender_label=to_gender_label(gender_code),
                language=str(row[idx["语言"]] or "").strip(),
                scene_l1=scene_l1,
                scene_l2=scene_l2,
                tech_desc=str(row[idx["技术描述（X系列&一句话）"]] or "").strip(),
                sort_value=sort_value,
            )
        )
    return records


def load_keyword_rows(xlsx_path: Path, taxonomy: TagTaxonomy) -> list[KeywordRow]:
    workbook = load_workbook(xlsx_path, data_only=True)
    rows = list(workbook[workbook.sheetnames[0]].iter_rows(values_only=True))
    results: list[KeywordRow] = []
    for row in rows[1:]:
        l1 = str(row[0] or "").strip()
        l2 = str(row[1] or "").strip()
        keywords: list[str] = []
        for cell in row[2:]:
            keywords.extend(split_tags(cell))
        if not l1 or not keywords:
            continue
        canonical_l1 = l1 if l1 in taxonomy.l1_tags else taxonomy.l2_to_l1.get(l2, l1)
        results.append(KeywordRow(scene_l1=canonical_l1, scene_l2=l2, keywords=keywords))
    return results


def build_augmented_keyword_rows(keyword_rows: list[KeywordRow], taxonomy: TagTaxonomy) -> list[KeywordRow]:
    merged: list[KeywordRow] = []
    seen: set[tuple[str, str, tuple[str, ...]]] = set()
    for row in keyword_rows:
        key = (row.scene_l1, row.scene_l2, tuple(sorted(set(row.keywords))))
        if key in seen:
            continue
        seen.add(key)
        merged.append(row)
    for l1 in sorted(taxonomy.l1_tags):
        merged.append(KeywordRow(scene_l1=l1, scene_l2="", keywords=[l1]))
    for l2 in sorted(taxonomy.l2_tags):
        merged.append(KeywordRow(scene_l1=taxonomy.l2_to_l1.get(l2, ""), scene_l2=l2, keywords=[l2]))
    return merged


def load_scene_keyword_overrides(rules: dict[str, Any], taxonomy: TagTaxonomy) -> list[KeywordRow]:
    rows: list[KeywordRow] = []
    for item in rules.get("scene_keyword_overrides", []):
        if not isinstance(item, dict):
            continue
        l1 = str(item.get("scene_l1", "") or "").strip()
        l2 = str(item.get("scene_l2", "") or "").strip()
        keywords = split_tags(item.get("keywords", ""))
        if not l1 or not keywords:
            continue
        canonical_l1 = l1 if l1 in taxonomy.l1_tags else taxonomy.l2_to_l1.get(l2, l1)
        rows.append(KeywordRow(scene_l1=canonical_l1, scene_l2=l2, keywords=keywords))
    return rows


def audio_choice_priority(row: dict[str, str], rules: dict[str, Any]) -> tuple[int, int, int]:
    name = row.get("speaker_name", "")
    text = row.get("audio_text", "")
    preferred_words = rules.get("audio_preferred_words", [])
    emotional_words = rules.get("audio_deprioritized_words", [])
    preferred = 0 if any(word in name or word in text for word in preferred_words) else 1
    emotional = 1 if any(word in name for word in emotional_words) else 0
    try:
        speaker_no = int(row.get("speaker_no", "999999999"))
    except ValueError:
        speaker_no = 999999999
    return (preferred, emotional, speaker_no)


def load_audio_url_map(csv_path: Path | None, rules: dict[str, Any]) -> dict[str, str]:
    if not csv_path or not csv_path.exists():
        return {}

    rows_by_vcn: dict[str, list[dict[str, str]]] = {}
    with csv_path.open("r", encoding="gb18030", newline="") as file:
        reader = csv.DictReader(file)
        for raw_row in reader:
            row = {key: (value or "").strip() for key, value in raw_row.items()}
            vcn = row.get("vcn", "")
            audio_url = row.get("audio_url", "")
            if not vcn or not audio_url:
                continue
            rows_by_vcn.setdefault(vcn, []).append(row)

    audio_map: dict[str, str] = {}
    for vcn, rows in rows_by_vcn.items():
        selected = sorted(rows, key=lambda row: audio_choice_priority(row, rules))[0]
        audio_map[vcn] = selected["audio_url"]
    return audio_map


def tech_rank(tech_desc: str, rules: dict[str, Any]) -> float:
    value = tech_desc or ""
    for item in rules.get("tech_rank_patterns", []):
        pattern = item.get("pattern")
        rank = item.get("rank")
        if pattern and rank is not None and re.search(pattern, value):
            return float(rank)
    match = re.search(r"[A-Za-z](\d+(?:\.\d+)?)", value)
    if match:
        return float(match.group(1))
    return 0.0


def tech_freshness_score(tech_desc: str, language: str, scene_l1: list[str], rules: dict[str, Any]) -> float:
    rank = tech_rank(tech_desc, rules)
    if rank <= 0:
        return 0.0
    excluded = set(rules.get("tech_freshness_excluded_l1", []))
    mainstream = bool(language) and not any(tag in excluded for tag in scene_l1)
    return (0.18 if mainstream else 0.08) * rank


def service_mismatch_penalty(scene_l1: list[str], has_service_hint: bool, rules: dict[str, Any]) -> float:
    if has_service_hint:
        return 0.0
    service_tags = set(rules.get("service_scene_tags", []))
    return 1.4 if any(tag in service_tags for tag in scene_l1) else 0.0


def over_broad_penalty(scene_l1: list[str], scene_l2: list[str]) -> float:
    return max(0, len(scene_l1) - 3) * 0.45 + max(0, len(scene_l2) - 6) * 0.18


def marketing_focus_score(voice: VoiceRecord, has_marketing_hint: bool, rules: dict[str, Any]) -> float:
    if not has_marketing_hint:
        return 0.0
    score = 0.0
    focus = rules.get("marketing_focus", {})
    l1_tags = set(focus.get("l1_tags", []))
    l2_tags = set(focus.get("l2_tags", []))
    short_video_tags = set(focus.get("short_video_tags", []))
    if any(tag in l1_tags for tag in voice.scene_l1) and len(voice.scene_l1) <= 3:
        score += 3.0
    if any(tag in l2_tags for tag in voice.scene_l2) and len(voice.scene_l2) <= 5:
        score += 2.0
    if any(tag in short_video_tags for tag in voice.scene_l1) and len(voice.scene_l1) <= 3:
        score += 1.0
    if len(voice.scene_l1) >= 6 or len(voice.scene_l2) >= 10:
        score -= 5.0
    return score


def speaker_penalty(voice: VoiceRecord, rules: dict[str, Any]) -> float:
    penalty = 0.0
    speaker_name = voice.speaker_name or ""
    normalized_name = normalize_text(speaker_name)
    vcn = voice.vcn or ""
    for item in rules.get("speaker_penalties", []):
        if not isinstance(item, dict):
            continue
        amount = float(item.get("penalty", 0.0) or 0.0)
        if amount <= 0:
            continue
        name_contains = str(item.get("name_contains", "") or "").strip()
        name_pattern = str(item.get("name_pattern", "") or "").strip()
        vcn_equals = str(item.get("vcn_equals", "") or "").strip()
        vcn_pattern = str(item.get("vcn_pattern", "") or "").strip()

        matched = False
        if name_contains and normalize_text(name_contains) in normalized_name:
            matched = True
        if name_pattern and re.search(name_pattern, speaker_name):
            matched = True
        if vcn_equals and vcn_equals == vcn:
            matched = True
        if vcn_pattern and re.search(vcn_pattern, vcn):
            matched = True
        if matched:
            penalty += amount
    return penalty


def language_match_score(voice_language: str, hit_info: dict[str, Any]) -> float:
    detected_language = str(hit_info.get("detected_language", "") or "").strip()
    lang_hits = set(hit_info.get("lang_hits", []))
    chinese_languages = {"普通话", "方言", "中文"}
    score = 0.0

    if detected_language:
        if voice_language == detected_language:
            score += 8.0
        elif voice_language in chinese_languages:
            score -= 8.0
        elif voice_language:
            score -= 2.0

    if lang_hits:
        if voice_language in lang_hits:
            score += 4.0
        elif voice_language:
            score -= 0.5
    return score


def build_keyword_hits(text: str, keyword_rows: list[KeywordRow], rules: dict[str, Any]) -> dict[str, Any]:
    normalized = normalize_text(text)
    l1_hits: Counter[str] = Counter()
    l2_hits: Counter[str] = Counter()
    matched_pairs: list[dict[str, Any]] = []
    for row in keyword_rows:
        matched = [keyword for keyword in row.keywords if keyword and keyword in normalized]
        if not matched:
            continue
        weighted_count = sum(keyword_weight(keyword, rules) for keyword in matched)
        if weighted_count <= 0:
            continue
        l1_hits[row.scene_l1] += weighted_count
        if row.scene_l2:
            l2_hits[row.scene_l2] += weighted_count
        matched_pairs.append(
            {
                "scene_l1": row.scene_l1,
                "scene_l2": row.scene_l2,
                "matched_keywords": matched[:10],
                "matched_count": len(matched),
                "matched_weighted_count": round(weighted_count, 4),
            }
        )

    intent_hints = rules.get("intent_hints", {})
    marketing_hints = intent_hints.get("marketing", [])
    service_hints = intent_hints.get("service", [])
    formal_hints = intent_hints.get("formal", [])
    attribute_hints = rules.get("attribute_hints", {})
    marketing_score = sum(1 for keyword in marketing_hints if keyword in normalized)
    price_pattern = rules.get("price_pattern", r"\d+(\.\d+)?")
    price_hit = bool(re.search(price_pattern, normalized))
    return {
        "l1_hits": l1_hits,
        "l2_hits": l2_hits,
        "matched_pairs": matched_pairs,
        "lang_hits": [lang for lang in rules.get("language_hints", []) if lang in normalized],
        "has_child_hint": any(keyword in normalized for keyword in attribute_hints.get("child", [])),
        "has_dialect_hint": any(keyword in normalized for keyword in attribute_hints.get("dialect", [])),
        "has_service_hint": any(keyword in normalized for keyword in service_hints),
        "has_marketing_hint": marketing_score >= 2 or (marketing_score >= 1 and price_hit),
        "marketing_score": marketing_score,
        "price_hit": price_hit,
        "formal_score": sum(1 for keyword in formal_hints if keyword in normalized),
        "scene_weight_total": float(sum(l1_hits.values()) + sum(l2_hits.values())),
    }


def compute_rule_score(voice: VoiceRecord, hit_info: dict[str, Any], rules: dict[str, Any]) -> tuple[float, dict[str, Any]]:
    l1_score = 0.0
    l2_score = 0.0
    attr_score = 0.0
    lang_score = 0.0
    matched_l1: list[str] = []
    matched_l2: list[str] = []

    for tag in voice.scene_l1:
        if tag in hit_info["l1_hits"]:
            l1_score += 3.0 * hit_info["l1_hits"][tag]
            matched_l1.append(tag)
    for tag in voice.scene_l2:
        if tag in hit_info["l2_hits"]:
            l2_score += 5.0 * hit_info["l2_hits"][tag]
            matched_l2.append(tag)

    child_tag = rules.get("attribute_tags", {}).get("child", "")
    dialect_tag = rules.get("attribute_tags", {}).get("dialect", "")
    lang_score += language_match_score(voice.language, hit_info)

    if child_tag and hit_info["has_child_hint"] and (child_tag in voice.scene_l1 or child_tag in voice.scene_l2):
        attr_score += 2.5
    if dialect_tag and hit_info["has_dialect_hint"] and (dialect_tag in voice.scene_l1 or dialect_tag in voice.scene_l2 or voice.language == dialect_tag):
        attr_score += 2.5

    tech_score = tech_freshness_score(voice.tech_desc, voice.language, voice.scene_l1, rules)
    marketing_score = marketing_focus_score(voice, hit_info["has_marketing_hint"], rules)
    broad_penalty = over_broad_penalty(voice.scene_l1, voice.scene_l2)
    service_penalty = service_mismatch_penalty(voice.scene_l1, hit_info["has_service_hint"], rules)
    speaker_penalty_score = speaker_penalty(voice, rules)
    tie_breaker = max(0.0, (1000.0 - min(voice.sort_value, 1000.0)) / 1000.0)
    score = l1_score + l2_score + attr_score + lang_score + tech_score + marketing_score + tie_breaker - broad_penalty - service_penalty - speaker_penalty_score
    return score, {
        "matched_l1": matched_l1,
        "matched_l2": matched_l2,
        "l1_score": l1_score,
        "l2_score": l2_score,
        "attr_score": attr_score,
        "lang_score": lang_score,
        "tech_score": tech_score,
        "marketing_focus_score": marketing_score,
        "broad_penalty": broad_penalty,
        "service_penalty": service_penalty,
        "speaker_penalty": speaker_penalty_score,
        "tie_breaker": tie_breaker,
    }


def voice_attributes(voice: VoiceRecord, rules: dict[str, Any]) -> list[str]:
    attrs: list[str] = []
    child_tag = rules.get("attribute_tags", {}).get("child", "")
    dialect_tag = rules.get("attribute_tags", {}).get("dialect", "")
    if child_tag and (child_tag in voice.scene_l1 or child_tag in voice.scene_l2):
        attrs.append(child_tag)
    if dialect_tag and (dialect_tag in voice.scene_l1 or dialect_tag in voice.scene_l2 or voice.language == dialect_tag):
        attrs.append(dialect_tag)
    return attrs


def build_matched_tags(item: dict[str, Any], hit_info: dict[str, Any]) -> dict[str, Any]:
    rule_reason = item.get("rule_reason", {})
    scene_l1 = [tag for tag in item.get("scene_l1", []) if tag in set(rule_reason.get("matched_l1", []))]
    scene_l2 = [tag for tag in item.get("scene_l2", []) if tag in set(rule_reason.get("matched_l2", []))]
    matched_attributes: list[str] = []
    if float(rule_reason.get("attr_score", 0.0) or 0.0) > 0:
        matched_attributes = list(item.get("attributes", []))
    language = item.get("language") == hit_info.get("detected_language") or (
        bool(hit_info.get("lang_hits")) and item.get("language") in set(hit_info.get("lang_hits", []))
    )
    return {
        "scene_l1": scene_l1,
        "scene_l2": scene_l2,
        "attributes": matched_attributes,
        "language": language,
    }


def speaker_dedupe_key(speaker_name: str, rules: dict[str, Any]) -> str:
    value = normalize_text(speaker_name)
    for pattern in rules.get("speaker_dedupe_suffix_patterns", []):
        value = re.sub(pattern, "", value).strip()
    for pattern in TECH_SUFFIX_PATTERNS:
        value = re.sub(pattern, "", value, flags=re.IGNORECASE).strip()
    return value or speaker_name


class VoiceRecommender:
    def __init__(
        self,
        *,
        voice_library_xlsx: Path,
        voice_taxonomy_xlsx: Path,
        voice_keywords_xlsx: Path,
        voice_audio_csv: Path | None,
        voice_rules_json: Path | None,
        deepseek_api_key: str,
        deepseek_api_base_url: str,
        deepseek_model: str,
        analysis_max_chars: int,
    ) -> None:
        self.rules = load_rule_settings(voice_rules_json)
        self.taxonomy = load_tag_taxonomy(voice_taxonomy_xlsx)
        self.voices = load_voice_records(voice_library_xlsx, self.taxonomy)
        self.available_languages = {voice.language for voice in self.voices if voice.language}
        keyword_rows = load_keyword_rows(voice_keywords_xlsx, self.taxonomy)
        self.keyword_rows = build_augmented_keyword_rows(
            keyword_rows + load_scene_keyword_overrides(self.rules, self.taxonomy),
            self.taxonomy,
        )
        self.audio_url_map = load_audio_url_map(voice_audio_csv, self.rules)
        self.deepseek_api_key = deepseek_api_key
        self.deepseek_api_base_url = deepseek_api_base_url
        self.deepseek_model = deepseek_model
        self.analysis_max_chars = analysis_max_chars

    def build_rule_candidates(self, text: str, candidate_size: int) -> dict[str, Any]:
        hit_info = build_keyword_hits(text, self.keyword_rows, self.rules)
        detected_language = detect_script_language(text, self.available_languages)
        hit_info["detected_language"] = detected_language
        if hit_info["has_marketing_hint"]:
            for item in self.rules.get("marketing_backfill", []):
                tag = item.get("tag")
                weight = float(item.get("weight", 0))
                if item.get("level") == "l1" and tag:
                    hit_info["l1_hits"][tag] += weight
                if item.get("level") == "l2" and tag:
                    hit_info["l2_hits"][tag] += weight
            hit_info["fallback_hint"] = "marketing_backfill"

        weak_scene = hit_info["scene_weight_total"] < 0.8
        if weak_scene:
            hit_info["l1_hits"] = Counter({key: value for key, value in hit_info["l1_hits"].items() if value >= 0.8})
            hit_info["l2_hits"] = Counter({key: value for key, value in hit_info["l2_hits"].items() if value >= 0.8})
            if hit_info["formal_score"] >= 2 and not hit_info["has_marketing_hint"]:
                for item in self.rules.get("formal_backfill", []):
                    tag = item.get("tag")
                    weight = float(item.get("weight", 0))
                    if item.get("level") == "l1" and tag:
                        hit_info["l1_hits"][tag] += weight
                    if item.get("level") == "l2" and tag:
                        hit_info["l2_hits"][tag] += weight
                hit_info["fallback_hint"] = "formal_narration_backfill"
            elif "fallback_hint" not in hit_info:
                hit_info["fallback_hint"] = "weak_scene_noisy_drop"
        elif "fallback_hint" not in hit_info:
            hit_info["fallback_hint"] = "none"

        scored = []
        for voice in self.voices:
            score, reason = compute_rule_score(voice, hit_info, self.rules)
            scored.append((voice, score, reason))
        scored.sort(key=lambda item: item[1], reverse=True)

        stage = "l2"
        filtered = [item for item in scored if item[2]["l2_score"] > 0]
        if len(filtered) < max(8, candidate_size // 3):
            stage = "l1"
            filtered = [item for item in scored if item[2]["l1_score"] > 0 or item[2]["l2_score"] > 0]
        if len(filtered) < max(8, candidate_size // 3):
            stage = "attr_lang"
            filtered = [item for item in scored if item[2]["attr_score"] > 0 or item[2]["lang_score"] > 0 or item[2]["l1_score"] > 0 or item[2]["l2_score"] > 0]
        if len(filtered) < candidate_size:
            if filtered:
                stage = f"{stage}_with_global_backfill"
                present = {item[0].vcn for item in filtered}
                filtered = filtered + [item for item in scored if item[0].vcn not in present]
            else:
                stage = "global_fallback"
                filtered = scored

        if detected_language:
            language_matched = [item for item in filtered if item[0].language == detected_language]
            if language_matched:
                stage = f"{stage}_language_strict"
                filtered = language_matched + [item for item in filtered if item[0].language != detected_language]

        candidates = []
        candidate_limit = len(language_matched) if detected_language and language_matched else candidate_size
        for index, (voice, score, reason) in enumerate(filtered[:candidate_limit]):
            candidates.append(
                {
                    "candidate_index": index,
                    "speaker_name": voice.speaker_name,
                    "vcn": voice.vcn,
                    "gender": voice.gender_label,
                    "language": voice.language,
                    "scene_l1": voice.scene_l1,
                    "scene_l2": voice.scene_l2,
                    "attributes": voice_attributes(voice, self.rules),
                    "tech_desc": voice.tech_desc,
                    "rule_score": round(score, 6),
                    "rule_reason": reason,
                }
            )

        return {
            "stage": stage,
            "hit_info": {
                "l1_hits": dict(hit_info["l1_hits"]),
                "l2_hits": dict(hit_info["l2_hits"]),
                "lang_hits": hit_info["lang_hits"],
                "detected_language": detected_language,
                "has_child_hint": hit_info["has_child_hint"],
                "has_dialect_hint": hit_info["has_dialect_hint"],
                "has_service_hint": hit_info["has_service_hint"],
                "has_marketing_hint": hit_info["has_marketing_hint"],
                "marketing_score": hit_info["marketing_score"],
                "price_hit": hit_info["price_hit"],
                "formal_score": hit_info["formal_score"],
                "scene_weight_total": round(hit_info["scene_weight_total"], 4),
                "fallback_hint": hit_info["fallback_hint"],
                "matched_pairs": hit_info["matched_pairs"][:12],
            },
            "candidates": candidates,
        }

    def recommend(self, *, text: str, top_k: int, candidate_size: int, use_deepseek: bool = True) -> dict[str, Any]:
        rule_pack = self.build_rule_candidates(text, candidate_size)
        llm_ranked = None
        llm_error = None
        llm_raw = ""
        mode = "rule_only"
        if use_deepseek and self.deepseek_api_key:
            try:
                response = rerank_with_deepseek(
                    api_key=self.deepseek_api_key,
                    api_base_url=self.deepseek_api_base_url,
                    model=self.deepseek_model,
                    timeout_seconds=30,
                    max_script_chars=self.analysis_max_chars,
                    script_text=text,
                    candidates=rule_pack["candidates"],
                )
                llm_ranked = response["ranked"]
                llm_raw = response["raw_content"]
                mode = "deepseek_rerank"
            except Exception as exc:
                llm_error = str(exc)
                mode = "rule_only_after_deepseek_error"

        results = self._merge_ranking(
            candidates=rule_pack["candidates"],
            llm_ranked=llm_ranked,
            top_k=top_k,
            rule_weight=0.7 if rule_pack["hit_info"]["has_marketing_hint"] else 0.4,
        )
        recommendations = [
            Recommendation(
                rank=rank + 1,
                speaker_name=item["speaker_name"],
                vcn=item["vcn"],
                gender=item["gender"],
                language=item["language"],
                scene_l1=item["scene_l1"],
                scene_l2=item["scene_l2"],
                attributes=item["attributes"],
                matched_tags=build_matched_tags(item, rule_pack["hit_info"]),
                tech_desc=item["tech_desc"],
                audio_url=self.audio_url_map.get(item["vcn"]),
                score=round(float(item["final_score"]), 6),
                reason=item["llm_reason"],
                debug={
                    "rule_score": item["rule_score"],
                    "llm_score": item["llm_score"],
                    "rule_reason": item["rule_reason"],
                    "candidate_index": item["candidate_index"],
                },
            )
            for rank, item in enumerate(results)
        ]
        return {
            "mode": mode,
            "stage": rule_pack["stage"],
            "hit_info": rule_pack["hit_info"],
            "recommendations": recommendations,
            "llm_error": llm_error,
            "llm_raw": llm_raw,
        }

    def _merge_ranking(self, candidates: list[dict[str, Any]], llm_ranked: list[dict[str, Any]] | None, top_k: int, rule_weight: float) -> list[dict[str, Any]]:
        by_index = {candidate["candidate_index"]: candidate for candidate in candidates}
        merged: list[dict[str, Any]] = []
        if llm_ranked:
            for item in llm_ranked:
                index = item.get("candidate_index")
                if index not in by_index:
                    continue
                candidate = dict(by_index[index])
                llm_score = float(item.get("score", 0.0))
                candidate["llm_score"] = llm_score
                candidate["llm_reason"] = str(item.get("reason", "")).strip()
                candidate["final_score"] = rule_weight * float(candidate["rule_score"]) + (1.0 - rule_weight) * llm_score
                merged.append(candidate)
        if not merged:
            for candidate in candidates:
                fallback = dict(candidate)
                fallback["llm_score"] = None
                fallback["llm_reason"] = "规则召回排序"
                fallback["final_score"] = float(fallback["rule_score"])
                merged.append(fallback)

        present = {item["candidate_index"] for item in merged}
        for candidate in candidates:
            if candidate["candidate_index"] in present:
                continue
            fallback = dict(candidate)
            fallback["llm_score"] = None
            fallback["llm_reason"] = "模型未返回该候选，使用规则补齐"
            fallback["final_score"] = float(fallback["rule_score"])
            merged.append(fallback)

        best_by_speaker: dict[str, dict[str, Any]] = {}
        for item in merged:
            dedupe_key = speaker_dedupe_key(item["speaker_name"], self.rules)
            current = best_by_speaker.get(dedupe_key)
            item["_dedupe_tech_rank"] = tech_rank(str(item.get("tech_desc", "")), self.rules)
            if current is None:
                best_by_speaker[dedupe_key] = item
                continue

            current_key = (
                float(current.get("_dedupe_tech_rank", 0.0)),
                float(current.get("final_score", 0.0)),
                -int(current.get("candidate_index", 999999)),
            )
            item_key = (
                float(item.get("_dedupe_tech_rank", 0.0)),
                float(item.get("final_score", 0.0)),
                -int(item.get("candidate_index", 999999)),
            )
            if item_key > current_key:
                best_by_speaker[dedupe_key] = item

        deduped = sorted(best_by_speaker.values(), key=lambda item: float(item["final_score"]), reverse=True)
        for item in deduped:
            item.pop("_dedupe_tech_rank", None)
        deduped = deduped[:top_k]
        return deduped


def recommendation_to_dict(item: Recommendation) -> dict[str, Any]:
    return {
        "rank": item.rank,
        "speaker_name": item.speaker_name,
        "vcn": item.vcn,
        "gender": item.gender,
        "language": item.language,
        "scene_l1": item.scene_l1,
        "scene_l2": item.scene_l2,
        "attributes": item.attributes,
        "matched_tags": item.matched_tags,
        "tech_desc": item.tech_desc,
        "audio_url": item.audio_url,
        "score": item.score,
        "reason": item.reason,
        "debug": item.debug,
    }


def dump_json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False)
